import os
import datetime
import logging
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl import load_workbook
import subprocess
import platform
import pandas as pd
from webdriver_manager.chrome import ChromeDriverManager
from selenium.webdriver.chrome.service import Service as ChromeService
from selenium.webdriver.chrome.options import Options
import re
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support.ui import Select
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import TimeoutException, UnexpectedAlertPresentException
import getpass
import ctypes
import time
from selenium.webdriver import ActionChains
import threading
import json


def split_username(username):
    import re
    # Try to split on common delimiters
    if "_" in username:
        parts = username.split("_")
    elif "." in username:
        parts = username.split(".")
    else:
        # Try to split camel case or just capitalize first and last
        parts = re.findall(r'[A-Z][a-z]*', username)
        if not parts:
            # Try to split in the middle for two names if all lowercase
            if len(username) > 6:
                mid = len(username) // 2
                parts = [username[:mid], username[mid:]]
            else:
                parts = [username]
    return " ".join([p.capitalize() for p in parts if p])


def get_full_username():
    try:
        GetUserNameEx = ctypes.windll.secur32.GetUserNameExW
        NameDisplay = 3

        size = ctypes.pointer(ctypes.c_ulong(0))
        GetUserNameEx(NameDisplay, None, size)

        name_buffer = ctypes.create_unicode_buffer(size.contents.value)
        GetUserNameEx(NameDisplay, name_buffer, size)

        return name_buffer.value
    except Exception:
        username = getpass.getuser()
        return split_username(username)


def clean_display_name(full_name):
    import re
    cleaned = re.sub(r"\(.*?\)", "", full_name)
    if "," in cleaned:
        cleaned = cleaned.split(",")[1]
    return cleaned.strip().title()


def safe_fill_field(self, xpath, value, field_name="Field"):
    try:
        element = self.driver.find_element(By.XPATH, xpath)

        if not element.is_enabled() or element.get_attribute("readonly") == "true":
            logging.info(f"Skipped {field_name} because it's not interactable.")
            return

        current_val = element.get_attribute("value")
        if current_val.strip() == str(value).strip():
            logging.info(f"{field_name} already set correctly. Skipping.")
            return

        # Clear and fill in one go
        element.clear()
        element.send_keys(value)
        logging.info(f"{field_name} set to: {value}")

        # Handle any popup that might have appeared immediately
        try:
            alert = self.driver.switch_to.alert
            alert_text = alert.text.strip()
            alert.accept()

            if "duplicate document" in alert_text.lower():
                logging.info(f"Duplicate alert detected after {field_name}. Handling...")
                self.handle_duplicate_lni_popup()
                # If this was a comments field, retry the fill
                if field_name == "Comments":
                    element.clear()
                    element.send_keys(value)
                    logging.info(f"Retried filling {field_name} after duplicate alert")
        except:
            pass  # No alert present, continue normally

    except Exception as e:
        logging.warning(f"Failed to set {field_name} at {xpath}")
        # Check if the error was due to a duplicate popup
        try:
            alert = self.driver.switch_to.alert
            alert_text = alert.text.strip()
            alert.accept()

            if "duplicate document" in alert_text.lower():
                logging.info("Duplicate alert detected during error handling. Processing...")
                self.handle_duplicate_lni_popup()
                # Retry the field fill
                element.clear()
                element.send_keys(value)
                logging.info(f"Retried filling {field_name} after duplicate alert")
        except:
            pass  # No alert present, continue with original error


# ========== Load Config ==========
CONFIG_DIR = Path(__file__).parent / "config"


def load_config():
    try:
        config_path = CONFIG_DIR / "config.json"
        if config_path.exists():
            with open(config_path, "r") as f:
                return json.load(f)
        return {"headless": False}
    except Exception as e:
        logging.error(f"Error loading config")
        return {"headless": False}


# ========== Global Error Log ==========
status_updates_buffer = {}
error_log_entries = []


# ========== Setup Logging ==========
def setup_logging():
    import sys
    from colorama import Fore, Style, init
    init(autoreset=True)

    class ColorFormatter(logging.Formatter):
        COLORS = {
            'INFO': Fore.GREEN,
            'WARNING': Fore.YELLOW,
            'ERROR': Fore.RED,
        }

        def format(self, record):
            color = self.COLORS.get(record.levelname, "")
            time_str = datetime.datetime.now().strftime("%I:%M:%S %p")
            message = f"{time_str} - {record.levelname} - {record.getMessage()}"
            return color + message + Style.RESET_ALL

    class PlainFormatter(logging.Formatter):
        def format(self, record):
            time_str = datetime.datetime.now().strftime("%I:%M:%S %p")
            return f"{time_str} - {record.levelname} - {record.getMessage()}"

    logs_path = Path.home() / "Downloads" / "Case Law Auto-Routing Resources" / "Logs"
    logs_path.mkdir(parents=True, exist_ok=True)

    now = datetime.datetime.now().strftime("%I-%M-%S_%p").lstrip("0")
    log_file = logs_path / f"log_{now}.txt"

    # Create console and file handlers separately
    file_handler = logging.FileHandler(log_file, encoding='utf-8')
    file_handler.setFormatter(PlainFormatter())

    console_handler = logging.StreamHandler(sys.stdout)
    console_handler.setFormatter(ColorFormatter())

    logging.basicConfig(
        level=logging.INFO,
        handlers=[file_handler, console_handler]
    )

    logging.info("Logging initialized.")


# ========== Folder Creation ==========
def create_folders():
    downloads_path = str(Path.home() / "Downloads")
    base_folder = os.path.join(downloads_path, "Case Law Auto-Routing Resources")
    logs_folder = os.path.join(base_folder, "Logs")
    os.makedirs(logs_folder, exist_ok=True)
    logging.info("Created folders at Downloads.")
    return base_folder


# ========== Excel File Creation ==========
def generate_excel(folder_path):
    now = datetime.datetime.now()
    time_str = now.strftime("%I%M%S_%p").lstrip("0")
    date_str = now.strftime("%d%m%Y")
    filename = f"Case Law Auto-Routing Data Mapping Sheet - {time_str} - {date_str}.xlsx"
    file_path = os.path.join(folder_path, filename)

    wb = Workbook()
    ws = wb.active
    ws.title = "Mapping Data"

    headers = [
        "Require Cardinal Process", "Court Code", "File Name", "LNI",
        "Tier", "Received Date", "Time Left",
        "Recycled Counsel LNI", "Source Detail", "Comments",
        "Decision Date",  # Optional manual decision date for this row
        "Status"
    ]

    # Style Definitions
    header_font = Font(name="Aptos Display", size=12, bold=True)
    alignment = Alignment(horizontal="center", vertical="center")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    fill = PatternFill(start_color="E7CCFC", end_color="E7CCFC", fill_type="solid")

    # Write Headers
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        cell.alignment = alignment
        cell.border = border
        cell.fill = fill

    # Format date columns (Received Date and Decision Date) for readability
    # Received Date: column F, Decision Date: column K
    for col_letter in ['F', 'K']:
        try:
            ws.column_dimensions[col_letter].number_format = 'MM/DD/YYYY'
        except Exception:
            pass

    # Autofit columns
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column].width = adjusted_width

    wb.save(file_path)
    logging.info(f"Excel file created: {file_path}")

    return file_path


# ========== Open Excel File ==========
def open_excel_file(path):
    try:
        if platform.system() == "Windows":
            os.startfile(path)
        elif platform.system() == "Darwin":
            subprocess.call(["open", path])
        else:
            subprocess.call(["xdg-open", path])
        logging.info("Excel file opened for user input.")
    except Exception as e:
        logging.error(f"Error opening file")


# ========== Get Latest Excel File ==========
def get_latest_excel_file(folder_path):
    files = [f for f in os.listdir(folder_path) if f.endswith(".xlsx")]
    if not files:
        logging.error("No Excel files found in the folder.")
        return None
    files.sort(key=lambda x: os.path.getmtime(os.path.join(folder_path, x)), reverse=True)
    latest_file = os.path.join(folder_path, files[0])
    logging.info(f"Latest Excel file detected: {latest_file}")
    return latest_file


# ========== Read Excel Mapping Data ==========
def read_mapping_data(file_path):
    """
    Read the Mapping Data sheet and return only the columns needed for routing.

    Column layout (1-based Excel columns):
      A: Require Cardinal Process
      B: Court Code
      C: File Name
      D: LNI
      E: Tier
      F: Received Date
      G: Time Left
      H: Recycled Counsel LNI
      I: Source Detail
      J: Comments
      K: Decision Date   (new manual override column)
      L: Status

    We keep Decision Date as its own column so it can act as the primary
    source for the IRT form's decision date when provided by the user.
    """
    try:
        df = pd.read_excel(file_path, sheet_name="Mapping Data", engine="openpyxl")

        # Select only the columns we care about, by index in the sheet
        selected_columns = df.iloc[:, [0, 2, 3, 5, 7, 8, 9, 10, 11]].copy()
        selected_columns.columns = [
            "Require Cardinal Process",
            "FileName",
            "LNI",
            "Received Date",
            "Recycled Counsel LNI",
            "Source Detail",
            "Comments",
            "Decision Date",
            "Status",
        ]
        return selected_columns
    except Exception as e:
        logging.error(f"Error reading Excel file: {e}")
        return None


def is_counsel(file_name, dar_mode=False):
    file_name_str = str(file_name).lower()
    
    # Check for counsel in the filename (case-insensitive)
    if 'counsel' in file_name_str:
        return True
    
    # If DAR mode is on, also check for DAR-specific counsel patterns
    if dar_mode:
        # Check for DAR pattern with counsel
        if re.search(r'dar.*counsel', file_name_str, re.IGNORECASE):
            return True
    
    return False


def get_related_counsel_lnis(main_docket, full_df, recycled_lni=None, dar_mode=False):
    related_lnis = []

    # First, get all counsel LNIs with matching docket numbers
    for _, row in full_df.iterrows():
        file_name = str(row["FileName"]).strip()
        if is_counsel(file_name, dar_mode):
            docket = extract_docket_number(file_name, dar_mode)
            if docket == main_docket:
                lni = str(row["LNI"]).strip()
                if lni.lower() != "nan":
                    related_lnis.append(lni)

    # Then handle recycled LNIs if present
    if recycled_lni and str(recycled_lni).strip().lower() != "nan":
        # Split by semicolon and handle each LNI
        recycled_lnis = [lni.strip() for lni in str(recycled_lni).split(';') if lni.strip()]
        for lni in recycled_lnis:
            if lni.lower() != "nan" and lni not in related_lnis:
                related_lnis.append(lni)

    return related_lnis


def extract_docket_number(file_name, dar_mode=False):
    """
    Extract base docket number from filename, supporting both SMD and DAR patterns.
    
    SMD Examples:
    - "LDC_SMD_24-7640a_E2E_PCQ.pdf" -> "24-7640"
    - "LDC_SMD_24-7640counsel_E2E.htm" -> "24-7640"
    
    DAR Examples:
    - "DAR_924cv80713-56_E2E.pdf" -> "24-80713"
    - "dar_2-23CV2039_PAWD_55_20250804_140135898.pdf" -> "23-2039"
    """
    if dar_mode:
        # Try DAR pattern first (case-insensitive)
        dar_match = re.search(r'dar_[\w\-]*(\d{2})[-_]?(cv|md|cd|mc|cr|mj)(\d+)', str(file_name), re.IGNORECASE)
        if dar_match:
            before_court_raw = dar_match.group(1)
            after_court = dar_match.group(3)
            
            # Extract the last 2 digits from before_court (the actual year)
            if len(before_court_raw) >= 2:
                year = before_court_raw[-2:]
                return f"{year}-{after_court}"
            else:
                return f"{before_court_raw}-{after_court}"
        
        # Also support LDC_PC / LDC_RR DAR-style patterns
        ldc_match = re.search(r'ldc_(pc|rr)_[\w\-]*(\d{2})[-_]?(cv|md|cd|mc|cr|mj)(\d+)', str(file_name), re.IGNORECASE)
        if ldc_match:
            before_court_raw = ldc_match.group(2)
            after_court = ldc_match.group(4)
            year = before_court_raw[-2:] if len(before_court_raw) >= 2 else before_court_raw
            return f"{year}-{after_court}"
    
    # Try SMD pattern (original)
    file_name = re.sub(r"counsel-\d+", "counsel", str(file_name))
    smd_match = re.search(r"LDC_SMD([_\-])([\d\-]+)[a-z]?", file_name)
    if smd_match:
        return smd_match.group(2)
    
    # If DAR mode is on but no DAR pattern found, try SMD as fallback
    if dar_mode:
        smd_match = re.search(r"LDC_SMD([_\-])([\d\-]+)[a-z]?", file_name)
        if smd_match:
            return smd_match.group(2)
    
    return None


def filter_mapping_data(df, dar_mode=False):
    df = df.copy()
    df["FileType"] = df["FileName"].apply(lambda x: "Counsel" if is_counsel(x, dar_mode) else "Main")

    counsel_df = df[df["FileType"] == "Counsel"].copy()
    main_df = df[df["FileType"] == "Main"].copy()

    logging.info(f"Counsel Count: {len(counsel_df)}, Main Opinion Count: {len(main_df)}")
    return counsel_df, main_df


def get_source_detail_mapping():
    """Return mapping of abbreviations to full source detail names"""
    return {
        'ao': 'Adopting Order',
        'cs': 'Change Status',
        'cd': 'Concur in Part / Dissent in Part',
        'cop': 'Concurring Opinion, Concur in Part',
        'cc': 'Corrections',
        'ci': 'Counsel Information',
        'dd': 'Disposition Only',
        'dop': 'Dissenting Opinion, Dissent in Part',
        'e': 'Excluded',
        'ex': 'Exhibits, Attachments, Appendix',
        'j': 'Judgement',
        'le': 'Letter',
        'lod': 'List of Documents',
        'm': 'Minutes',
        'op': 'Opinion',
        'or': 'Order',
        'ornol': 'Order (Not Online)',
        'orol': 'Order (Online)',
        'pool': 'Published Opinion (Online)',
        'rl': 'Release',
        'rr': 'Reports and Recommendations/Finding of Facts/Conclusion of Law',
        's': 'Settlement',
        'sc': 'Scheduling',
        'so': 'Standing Order',
        'sp': 'Special Proceeding',
        'st': 'Status',
        'su': 'Subpoena',
        'sw': 'Sworn Statement',
        't': 'Transcript',
        'w': 'Warrant',
        'wa': 'Waiver',
        'wo': 'Writ of Order',
        'wp': 'Written Pleading',
        'ws': 'Written Statement',
        'wt': 'Written Testimony'
    }


def resolve_source_detail(abbreviation):
    """Resolve source detail abbreviation to full name"""
    mapping = get_source_detail_mapping()
    return mapping.get(abbreviation.lower(), abbreviation)


def flush_status_updates(file_path, status_map):
    try:
        wb = load_workbook(file_path)
        ws = wb["Mapping Data"]

        # Define status styles
        status_styles = {
            "DONE": {"bold": True, "color": "00A86B"},
            "ALREADY PROCESSED": {"bold": True, "color": "4F39BC"},
            "ERROR: LNI NOT FOUND": {"bold": True, "color": "E3242B"},
            "ERROR: INVALID LNI FORMAT": {"bold": True, "color": "E3242B"},
            "ERROR: UNABLE TO LOAD IRT FORM WINDOW": {"bold": True, "color": "E3242B"},
            "ERROR": {"bold": True, "color": "E3242B"},
            "RELATED LNI ERROR": {"bold": True, "color": "E3242B"},
            "NO COUNSEL ATTACHED": {"bold": True, "color": "E3242B"}
        }

        for row_index, status in status_map.items():
            # Status now lives in column L (12); column K is Decision Date
            cell = ws.cell(row=row_index + 2, column=12)  # Column L
            cell.value = status

            # Apply formatting based on status
            style = status_styles.get(status, {"bold": False, "color": "000000"})
            cell.font = Font(bold=style["bold"], color=style["color"])

        # Autofit all columns
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column].width = adjusted_width

        wb.save(file_path)
        logging.info(f"Status updates flushed to Excel file: {file_path}")

    except Exception as e:
        logging.error(f"Error flushing status updates: {e}")
